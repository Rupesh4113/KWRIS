"""Tests for Report Generation and Formula Injection Sanitization."""

import pytest
import io
import pandas as pd
import openpyxl
from src.reporting import sanitize_cell_value, sanitize_dataframe, ReportGenerator


def test_formula_injection_sanitization():
    # Attack strings that attempt formula injection
    dangerous_inputs = ["=cmd|' /C calc'!A0", "+1+1", "-2+3", "@SUM(A1:A10)", "\tmalicious"]
    for val in dangerous_inputs:
        safe_val = sanitize_cell_value(val)
        assert safe_val.startswith("'")


def test_excel_report_generation():
    generator = ReportGenerator()
    summary = {
        "filename": "test_data.csv",
        "total_records": 10,
        "total_stations": 2,
        "quality_score": 92.5
    }
    issues_df = pd.DataFrame([{
        "dataset_name": "test_data.csv",
        "station_id": "STN_A",
        "observation_date": "2024-06-01",
        "original_value": "-5.0",
        "issue_type": "Negative Rainfall",
        "severity": "Critical",
        "validation_rule": "V_NEGATIVE_VALUE",
        "explanation": "Negative value",
        "recommended_action": "Set to 0",
        "review_status": "Pending"
    }])
    priorities_df = pd.DataFrame([{
        "station_id": "STN_A",
        "priority_category": "Critical",
        "priority_score": 85.0
    }])
    predictions_df = pd.DataFrame([{
        "station_id": "STN_A",
        "date": "2024-06-02",
        "predicted_value_mm": 12.5,
        "is_estimated": True
    }])
    audit_df = pd.DataFrame([{
        "station_id": "STN_A",
        "action": "Accepted"
    }])
    dataset_df = pd.DataFrame([{
        "station_id": "STN_A",
        "date": "2024-06-01",
        "rainfall": 10.0
    }])

    excel_buffer = generator.generate_excel_report(
        summary, issues_df, priorities_df, predictions_df, audit_df, dataset_df
    )
    assert excel_buffer is not None
    excel_buffer.seek(0)
    wb = openpyxl.load_workbook(excel_buffer)
    assert "Executive_Summary" in wb.sheetnames
    assert "Data_Issues" in wb.sheetnames
    assert "Station_Priority_Health" in wb.sheetnames
    assert "Missing_Predictions" in wb.sheetnames
